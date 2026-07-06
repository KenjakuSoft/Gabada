import csv
import io
import logging
import re
from typing import Any, Optional
from urllib.parse import parse_qsl, quote
from urllib.parse import urlparse, urlunparse

import httpx
import openpyxl
import pdfplumber
from fastapi import APIRouter, File, HTTPException, UploadFile

try:
    import xlrd
except Exception:  # pragma: no cover - optional dependency during local setup
    xlrd = None

from api.product_service import create_product_with_store, create_store_link
from database.db import get_db
from database.models import (
    ApiImportRequest,
    ApiPreviewRequest,
    ApiProductPreview,
    FileImportCommitRequest,
    FileImportCommitResponse,
    FileImportPreviewResponse,
    FileImportRow,
    ScrapeConfig,
    Store,
    WooImportRequest,
    WooProductPreview,
    WooPreviewRequest,
    WooSyncResult,
    neutralize_web_url,
)
from services.api_catalog import (
    fetch_api_preview_products,
    import_api_products,
    sync_all_enabled_api_stores,
    sync_api_product_store,
    sync_store_api_catalog,
)

router = APIRouter(prefix="/import", tags=["import"])
logger = logging.getLogger(__name__)

_WOO_API_MAX_PER_PAGE = 100
_WOO_MAX_PREVIEW_PAGES = 500
_WOO_TIMEOUT = httpx.Timeout(60.0, connect=20.0)
_KEEP_IMPORT_VALUE = object()


def _map_woo_stock(status: str) -> str:
    return {
        "instock": "in_stock",
        "outofstock": "out_of_stock",
        "onbackorder": "limited",
    }.get(status, "unknown")


def _woo_products_url(site_url: str) -> str:
    return f"{site_url.rstrip('/')}/wp-json/wc/v3/products"


def _woo_preview_error_detail(exc: Exception) -> str:
    if isinstance(exc, HTTPException) and isinstance(exc.detail, str):
        return exc.detail
    return str(exc)


def _clip_woo_body(text: str, *, limit: int = 220) -> str:
    collapsed = " ".join((text or "").split())
    if len(collapsed) <= limit:
        return collapsed
    return collapsed[: limit - 3].rstrip() + "..."


def _normalize_woo_site_url(site_url: str) -> str:
    raw = (site_url or "").strip()
    if not raw:
        raise HTTPException(status_code=400, detail="Site URL is required")
    if not re.match(r"^https?://", raw, re.I):
        raw = f"https://{raw}"
    parsed = urlparse(raw)
    if not parsed.netloc:
        raise HTTPException(status_code=400, detail="Enter a valid WooCommerce site URL")
    cleaned_path = parsed.path.rstrip("/")
    for suffix in ("/wp-json/wc/v3/products", "/wp-json/wc/v3"):
        if cleaned_path.endswith(suffix):
            cleaned_path = cleaned_path[: -len(suffix)]
            break
    if cleaned_path.endswith("/index.php"):
        cleaned_path = cleaned_path[: -len("/index.php")]
    return urlunparse(
        (
            parsed.scheme.lower(),
            parsed.netloc,
            cleaned_path,
            "",
            "",
            "",
        )
    ).rstrip("/")


def _woo_products_urls(site_url: str) -> list[str]:
    base = _normalize_woo_site_url(site_url)
    candidates = [
        f"{base}/wp-json/wc/v3/products",
        f"{base}/index.php?rest_route=/wc/v3/products",
        f"{base}/?rest_route=/wc/v3/products",
    ]
    unique: list[str] = []
    seen: set[str] = set()
    for candidate in candidates:
        if candidate not in seen:
            unique.append(candidate)
            seen.add(candidate)
    return unique


def _woo_store_products_urls(site_url: str) -> list[str]:
    base = _normalize_woo_site_url(site_url)
    candidates = [
        f"{base}/wp-json/wc/store/v1/products",
        f"{base}/index.php?rest_route=/wc/store/v1/products",
        f"{base}/?rest_route=/wc/store/v1/products",
    ]
    unique: list[str] = []
    seen: set[str] = set()
    for candidate in candidates:
        if candidate not in seen:
            unique.append(candidate)
            seen.add(candidate)
    return unique


def _woo_per_page(value: int) -> int:
    return max(1, min(value, _WOO_API_MAX_PER_PAGE))


def _chunked(values: list[int], size: int) -> list[list[int]]:
    return [values[index : index + size] for index in range(0, len(values), size)]


def _woo_total_pages(resp: httpx.Response) -> int | None:
    raw = resp.headers.get("X-WP-TotalPages")
    if not raw:
        return None
    try:
        return max(1, int(raw))
    except ValueError:
        return None


def _woo_preview_item(product: dict) -> WooProductPreview:
    images = product.get("images", [])
    return WooProductPreview(
        woo_id=product["id"],
        name=product["name"],
        permalink=product["permalink"],
        regular_price=product.get("regular_price") or None,
        sale_price=product.get("sale_price") or None,
        stock_status=_map_woo_stock(product.get("stock_status", "outofstock")),
        image_url=images[0]["src"] if images else None,
        categories=[category["name"] for category in product.get("categories", [])],
    )


def _woo_is_public_product(product: dict) -> bool:
    return (product.get("status") or "").strip().lower() == "publish"


def _woo_store_product_stock_status(product: dict) -> str:
    if product.get("is_on_backorder"):
        return "onbackorder"

    stock_class = str((product.get("stock_availability") or {}).get("class") or "").strip().lower()
    if stock_class in {"in-stock", "instock"}:
        return "instock"
    if stock_class in {"out-of-stock", "outofstock"}:
        return "outofstock"

    if product.get("is_in_stock") is True:
        return "instock"
    if product.get("is_in_stock") is False:
        return "outofstock"
    return "outofstock"


def _woo_store_product_to_woo_shape(product: dict) -> dict:
    prices = product.get("prices") or {}
    regular_price = prices.get("regular_price") or prices.get("price") or None
    sale_price = prices.get("sale_price") or None
    if sale_price and regular_price and sale_price == regular_price:
        sale_price = None

    return {
        "id": product["id"],
        "name": product.get("name") or "",
        "permalink": product.get("permalink") or "",
        "regular_price": regular_price,
        "sale_price": sale_price,
        "stock_status": _woo_store_product_stock_status(product),
        "images": [
            {"src": image.get("src")}
            for image in product.get("images", [])
            if image.get("src")
        ],
        "categories": [
            {"name": category.get("name")}
            for category in product.get("categories", [])
            if category.get("name")
        ],
        "status": "publish",
    }


def _woo_request_target(endpoint: str, params: dict[str, Any]) -> tuple[str, dict[str, Any]]:
    parsed = urlparse(endpoint)
    merged_params = dict(parse_qsl(parsed.query, keep_blank_values=True))
    merged_params.update(params)
    request_url = urlunparse(parsed._replace(query=""))
    return request_url, merged_params


async def _woo_get_products(
    client: httpx.AsyncClient,
    *,
    site_url: str,
    consumer_key: str,
    consumer_secret: str,
    params: dict,
    preferred_endpoint: Optional[str] = None,
    preferred_auth_mode: Optional[str] = None,
) -> tuple[httpx.Response, str, str]:
    attempts: list[str] = []
    last_request_error: Optional[httpx.RequestError] = None
    endpoints = _woo_products_urls(site_url)
    if preferred_endpoint in endpoints:
        endpoints = [preferred_endpoint] + [
            endpoint for endpoint in endpoints if endpoint != preferred_endpoint
        ]

    auth_modes: list[tuple[str, dict[str, Any]]] = [
        (
            "basic auth",
            {
                "auth": (consumer_key, consumer_secret),
                "params": params,
            },
        ),
        (
            "query auth",
            {
                "auth": None,
                "params": {
                    **params,
                    "consumer_key": consumer_key,
                    "consumer_secret": consumer_secret,
                },
            },
        ),
    ]
    if preferred_auth_mode:
        auth_modes.sort(key=lambda item: 0 if item[0] == preferred_auth_mode else 1)

    for endpoint in endpoints:
        for label, request_kwargs in auth_modes:
            request_url, request_params = _woo_request_target(
                endpoint,
                request_kwargs["params"],
            )
            try:
                resp = await client.get(
                    request_url,
                    auth=request_kwargs["auth"],
                    params=request_params,
                )
            except httpx.RequestError as exc:
                last_request_error = exc
                attempts.append(f"{label} via {endpoint} unreachable: {exc}")
                continue

            if resp.status_code == 200:
                try:
                    resp.json()
                except ValueError:
                    attempts.append(f"{label} via {endpoint}: 200 OK but returned invalid JSON")
                    continue
                return resp, endpoint, label

            body_detail = _clip_woo_body(resp.text)
            status_detail = f"{resp.status_code} {resp.reason_phrase}".strip()
            if body_detail:
                attempts.append(f"{label} via {endpoint}: {status_detail} - {body_detail}")
            else:
                attempts.append(f"{label} via {endpoint}: {status_detail}")

    if last_request_error is not None and not attempts:
        raise HTTPException(status_code=502, detail=f"Cannot reach WooCommerce: {last_request_error}")

    detail = "WooCommerce API preview failed. " + " | ".join(attempts[:4])
    raise HTTPException(status_code=502, detail=detail)


async def _woo_get_public_store_products(
    client: httpx.AsyncClient,
    *,
    site_url: str,
    params: dict,
    preferred_endpoint: Optional[str] = None,
) -> tuple[httpx.Response, str]:
    attempts: list[str] = []
    last_request_error: Optional[httpx.RequestError] = None
    endpoints = _woo_store_products_urls(site_url)
    if preferred_endpoint in endpoints:
        endpoints = [preferred_endpoint] + [
            endpoint for endpoint in endpoints if endpoint != preferred_endpoint
        ]

    for endpoint in endpoints:
        request_url, request_params = _woo_request_target(endpoint, params)
        try:
            resp = await client.get(request_url, params=request_params)
        except httpx.RequestError as exc:
            last_request_error = exc
            attempts.append(f"public store api via {endpoint} unreachable: {exc}")
            continue

        if resp.status_code == 200:
            try:
                resp.json()
            except ValueError:
                attempts.append(f"public store api via {endpoint}: 200 OK but returned invalid JSON")
                continue
            return resp, endpoint

        body_detail = _clip_woo_body(resp.text)
        status_detail = f"{resp.status_code} {resp.reason_phrase}".strip()
        if body_detail:
            attempts.append(f"public store api via {endpoint}: {status_detail} - {body_detail}")
        else:
            attempts.append(f"public store api via {endpoint}: {status_detail}")

    if last_request_error is not None and not attempts:
        raise HTTPException(status_code=502, detail=f"Cannot reach WooCommerce public store API: {last_request_error}")

    detail = "WooCommerce public store API failed. " + " | ".join(attempts[:4])
    raise HTTPException(status_code=502, detail=detail)


async def _woo_fetch_preview_products_v3(body: WooPreviewRequest) -> list[WooProductPreview]:
    per_page = _woo_per_page(body.per_page)
    async with httpx.AsyncClient(timeout=_WOO_TIMEOUT) as client:
        resolved_endpoint: Optional[str] = None
        resolved_auth_mode: Optional[str] = None
        if not body.fetch_all:
            resp, _, _ = await _woo_get_products(
                client,
                site_url=body.credentials.site_url,
                consumer_key=body.credentials.consumer_key,
                consumer_secret=body.credentials.consumer_secret,
                params={
                    "per_page": per_page,
                    "page": max(1, body.page),
                    "status": "publish",
                },
                preferred_endpoint=resolved_endpoint,
                preferred_auth_mode=resolved_auth_mode,
            )
            return [
                _woo_preview_item(product)
                for product in resp.json()
                if _woo_is_public_product(product)
            ]

        previews: list[WooProductPreview] = []
        seen_ids: set[int] = set()
        repeated_pages = 0
        page = 1
        total_pages: int | None = None

        while True:
            if page > _WOO_MAX_PREVIEW_PAGES:
                raise HTTPException(
                    status_code=502,
                    detail="WooCommerce preview exceeded the safe page limit",
                )
            resp, resolved_endpoint, resolved_auth_mode = await _woo_get_products(
                client,
                site_url=body.credentials.site_url,
                consumer_key=body.credentials.consumer_key,
                consumer_secret=body.credentials.consumer_secret,
                params={
                    "per_page": per_page,
                    "page": page,
                    "status": "publish",
                },
                preferred_endpoint=resolved_endpoint,
                preferred_auth_mode=resolved_auth_mode,
            )
            products = [
                product for product in resp.json() if _woo_is_public_product(product)
            ]
            new_products = [
                product
                for product in products
                if product.get("id") not in seen_ids
            ]
            if new_products:
                previews.extend(_woo_preview_item(product) for product in new_products)
                seen_ids.update(product["id"] for product in new_products)
                repeated_pages = 0
            else:
                repeated_pages += 1

            if total_pages is None:
                total_pages = _woo_total_pages(resp)
            if total_pages is not None and page >= total_pages:
                break
            if len(products) < per_page:
                break
            if repeated_pages >= 1:
                break
            page += 1

        return previews


async def _woo_fetch_preview_products_store_api(body: WooPreviewRequest) -> list[WooProductPreview]:
    per_page = _woo_per_page(body.per_page)
    async with httpx.AsyncClient(timeout=_WOO_TIMEOUT) as client:
        resolved_endpoint: Optional[str] = None
        if not body.fetch_all:
            resp, _ = await _woo_get_public_store_products(
                client,
                site_url=body.credentials.site_url,
                params={
                    "per_page": per_page,
                    "page": max(1, body.page),
                },
                preferred_endpoint=resolved_endpoint,
            )
            return [
                _woo_preview_item(_woo_store_product_to_woo_shape(product))
                for product in resp.json()
            ]

        previews: list[WooProductPreview] = []
        seen_ids: set[int] = set()
        repeated_pages = 0
        page = 1
        total_pages: int | None = None

        while True:
            if page > _WOO_MAX_PREVIEW_PAGES:
                raise HTTPException(
                    status_code=502,
                    detail="WooCommerce public store preview exceeded the safe page limit",
                )
            resp, resolved_endpoint = await _woo_get_public_store_products(
                client,
                site_url=body.credentials.site_url,
                params={
                    "per_page": per_page,
                    "page": page,
                },
                preferred_endpoint=resolved_endpoint,
            )
            products = resp.json()
            new_products = [
                product
                for product in products
                if product.get("id") not in seen_ids
            ]
            if new_products:
                previews.extend(
                    _woo_preview_item(_woo_store_product_to_woo_shape(product))
                    for product in new_products
                )
                seen_ids.update(product["id"] for product in new_products)
                repeated_pages = 0
            else:
                repeated_pages += 1

            if total_pages is None:
                total_pages = _woo_total_pages(resp)
            if total_pages is not None and page >= total_pages:
                break
            if len(products) < per_page:
                break
            if repeated_pages >= 1:
                break
            page += 1

        return previews


async def _woo_fetch_preview_products(body: WooPreviewRequest) -> list[WooProductPreview]:
    try:
        return await _woo_fetch_preview_products_v3(body)
    except HTTPException as exc:
        primary_detail = _woo_preview_error_detail(exc)
        logger.info(
            "Woo v3 preview failed for %s, trying public store API fallback: %s",
            body.credentials.site_url,
            primary_detail,
        )
        try:
            return await _woo_fetch_preview_products_store_api(body)
        except HTTPException as fallback_exc:
            fallback_detail = _woo_preview_error_detail(fallback_exc)
            raise HTTPException(
                status_code=502,
                detail=f"{primary_detail} | Public Store API fallback failed. {fallback_detail}",
            ) from fallback_exc


async def _woo_validate_preview_request(body: WooPreviewRequest) -> None:
    if not body.credentials.consumer_key.strip() or not body.credentials.consumer_secret.strip():
        raise HTTPException(status_code=400, detail="Consumer key and secret are required")
    _normalize_woo_site_url(body.credentials.site_url)


async def _woo_validate_import_request(body: WooImportRequest) -> None:
    if not body.credentials.consumer_key.strip() or not body.credentials.consumer_secret.strip():
        raise HTTPException(status_code=400, detail="Consumer key and secret are required")
    _normalize_woo_site_url(body.credentials.site_url)


def _normalize_store_product_url(url: str) -> str:
    text = (url or "").strip()
    if not text:
        return ""
    parsed = urlparse(text)
    path = parsed.path.rstrip("/")
    return urlunparse(
        (
            parsed.scheme.lower(),
            parsed.netloc.lower(),
            path,
            "",
            parsed.query,
            "",
        )
    )


def _normalize_store_product_name(name: str) -> str:
    return re.sub(r"\s+", " ", (name or "").strip()).lower()


def _row_to_store(row) -> Store:
    return Store(
        id=row["id"],
        name=row["name"],
        base_url=row["base_url"],
        manual_only=bool(row["manual_only"]),
        login_url=row["login_url"],
        storage_state_path=row["storage_state_path"],
        session_updated_at=row["session_updated_at"],
        woo_enabled=bool(row["woo_enabled"]),
        woo_site_url=row["woo_site_url"],
        woo_consumer_key=row["woo_consumer_key"],
        woo_consumer_secret=row["woo_consumer_secret"],
        woo_auto_import_new=bool(row["woo_auto_import_new"]),
        woo_last_synced_at=row["woo_last_synced_at"],
        created_at=row["created_at"],
    )


async def _get_store_or_404(db, store_id: int) -> Store:
    rows = await db.execute_fetchall(
        """SELECT
               id,
               name,
               base_url,
               manual_only,
               login_url,
               storage_state_path,
               session_updated_at,
               woo_enabled,
               woo_site_url,
               woo_consumer_key,
               woo_consumer_secret,
               woo_auto_import_new,
               woo_last_synced_at,
               created_at
           FROM stores
           WHERE id = ?""",
        (store_id,),
    )
    if not rows:
        raise HTTPException(status_code=404, detail="Store not found")
    return _row_to_store(rows[0])


def _woo_credentials_from_store(store: Store) -> WooPreviewRequest:
    if not store.woo_enabled:
        raise HTTPException(status_code=400, detail="Woo sync is not enabled for this store")
    if not store.woo_site_url or not store.woo_consumer_key or not store.woo_consumer_secret:
        raise HTTPException(status_code=400, detail="Store is missing Woo sync credentials")
    return WooPreviewRequest(
        credentials={
            "site_url": store.woo_site_url,
            "consumer_key": store.woo_consumer_key,
            "consumer_secret": store.woo_consumer_secret,
        },
        per_page=_WOO_API_MAX_PER_PAGE,
        page=1,
        fetch_all=True,
    )


async def _existing_store_product_maps(db, store_id: int) -> tuple[dict[str, dict], dict[str, dict]]:
    rows = await db.execute_fetchall(
        """SELECT
               p.id AS product_id,
               p.canonical_name,
               p.image_url,
               ps.id AS product_store_id,
               ps.product_url
           FROM product_stores ps
           JOIN products p ON p.id = ps.product_id
           WHERE ps.store_id = ?""",
        (store_id,),
    )
    by_url: dict[str, dict] = {}
    by_name: dict[str, dict] = {}
    for row in rows:
        normalized_url = _normalize_store_product_url(row["product_url"])
        normalized_name = _normalize_store_product_name(row["canonical_name"])
        payload = dict(row)
        if normalized_url and normalized_url not in by_url:
            by_url[normalized_url] = payload
        if normalized_name and normalized_name not in by_name:
            by_name[normalized_name] = payload
    return by_url, by_name


async def _find_any_product_by_name(db, canonical_name: str) -> Optional[dict]:
    rows = await db.execute_fetchall(
        """SELECT id, canonical_name, image_url
           FROM products
           WHERE lower(trim(canonical_name)) = lower(trim(?))
           ORDER BY id
           LIMIT 1""",
        (canonical_name,),
    )
    return dict(rows[0]) if rows else None


async def _merge_woo_catalog_item(
    db,
    *,
    store_id: int,
    name: str,
    permalink: str,
    image_url: Optional[str],
    existing_by_url: dict[str, dict],
    existing_by_name: dict[str, dict],
    create_if_missing: bool,
) -> str:
    # Woo permalinks/images come from the remote catalog, not the user —
    # drop non-web schemes instead of persisting them into <a href> data.
    permalink = neutralize_web_url(permalink) or ""
    image_url = neutralize_web_url(image_url)
    normalized_url = _normalize_store_product_url(permalink)
    normalized_name = _normalize_store_product_name(name)
    match = existing_by_url.get(normalized_url) or existing_by_name.get(normalized_name)
    if match is not None:
        if normalized_url and match.get("product_url") != permalink:
            await db.execute(
                "UPDATE product_stores SET product_url = ?, is_active = 1 WHERE id = ?",
                (permalink, match["product_store_id"]),
            )
            match["product_url"] = permalink
        if image_url and not match.get("image_url"):
            await db.execute(
                "UPDATE products SET image_url = ? WHERE id = ?",
                (image_url, match["product_id"]),
            )
            match["image_url"] = image_url
        if normalized_url:
            existing_by_url[normalized_url] = match
        if normalized_name:
            existing_by_name[normalized_name] = match
        return "existing"

    if not create_if_missing:
        return "skipped"

    scrape_config_json = ScrapeConfig(platform="woocommerce").model_dump_json()
    existing_product = await _find_any_product_by_name(db, name)
    if existing_product is not None:
        product_id = existing_product["id"]
        product_store_id = await create_store_link(
            db,
            product_id=product_id,
            store_id=store_id,
            product_url=permalink,
            scrape_config_json=scrape_config_json,
            is_active=True,
        )
        if image_url and not existing_product.get("image_url"):
            await db.execute(
                "UPDATE products SET image_url = ? WHERE id = ?",
                (image_url, product_id),
            )
    else:
        product_id, product_store_id = await create_product_with_store(
            db,
            canonical_name=name,
            image_url=image_url,
            category_id=None,
            store_id=store_id,
            product_url=permalink,
            scrape_config_json=scrape_config_json,
        )

    payload = {
        "product_id": product_id,
        "canonical_name": name,
        "image_url": image_url,
        "product_store_id": product_store_id,
        "product_url": permalink,
    }
    if normalized_url:
        existing_by_url[normalized_url] = payload
    if normalized_name:
        existing_by_name[normalized_name] = payload
    return "created"


async def sync_store_woo_catalog(store_id: int) -> WooSyncResult:
    db = await get_db()
    try:
        store = await _get_store_or_404(db, store_id)
        preview_request = _woo_credentials_from_store(store)
        await _woo_validate_preview_request(preview_request)
        previews = await _woo_fetch_preview_products(preview_request)
        existing_by_url, existing_by_name = await _existing_store_product_maps(db, store_id)

        created = 0
        existing = 0
        auto_import_new = bool(store.woo_auto_import_new)
        for preview in previews:
            outcome = await _merge_woo_catalog_item(
                db,
                store_id=store_id,
                name=preview.name,
                permalink=preview.permalink,
                image_url=preview.image_url,
                existing_by_url=existing_by_url,
                existing_by_name=existing_by_name,
                create_if_missing=auto_import_new,
            )
            if outcome == "created":
                created += 1
            elif outcome == "existing":
                existing += 1

        await db.execute(
            "UPDATE stores SET woo_last_synced_at = CURRENT_TIMESTAMP WHERE id = ?",
            (store_id,),
        )
        await db.commit()
        refreshed = await _get_store_or_404(db, store_id)
        return WooSyncResult(
            store_id=store_id,
            store_name=store.name,
            discovered=len(previews),
            created=created,
            existing=existing,
            auto_import_new=auto_import_new,
            synced_at=refreshed.woo_last_synced_at,
        )
    finally:
        await db.close()


async def sync_all_enabled_woo_stores() -> list[WooSyncResult]:
    db = await get_db()
    try:
        rows = await db.execute_fetchall(
            """SELECT id
               FROM stores
               WHERE woo_enabled = 1
                 AND woo_site_url IS NOT NULL
                 AND woo_consumer_key IS NOT NULL
                 AND woo_consumer_secret IS NOT NULL
               ORDER BY name"""
        )
        store_ids = [row["id"] for row in rows]
    finally:
        await db.close()

    results: list[WooSyncResult] = []
    for store_id in store_ids:
        try:
            results.append(await sync_store_woo_catalog(store_id))
        except Exception as exc:
            logger.error("Woo sync failed for store_id=%d: %s", store_id, exc)
    return results


async def _woo_fetch_products_by_ids_v3(body: WooImportRequest) -> list[dict]:
    product_ids = list(dict.fromkeys(body.product_ids))
    if not product_ids:
        return []

    products_by_id: dict[int, dict] = {}
    async with httpx.AsyncClient(timeout=_WOO_TIMEOUT) as client:
        resolved_endpoint: Optional[str] = None
        resolved_auth_mode: Optional[str] = None
        for chunk in _chunked(product_ids, _WOO_API_MAX_PER_PAGE):
            resp, resolved_endpoint, resolved_auth_mode = await _woo_get_products(
                client,
                site_url=body.credentials.site_url,
                consumer_key=body.credentials.consumer_key,
                consumer_secret=body.credentials.consumer_secret,
                params={
                    "include": ",".join(str(product_id) for product_id in chunk),
                    "per_page": len(chunk),
                    "status": "publish",
                },
                preferred_endpoint=resolved_endpoint,
                preferred_auth_mode=resolved_auth_mode,
            )
            for product in resp.json():
                if not _woo_is_public_product(product):
                    continue
                products_by_id[product["id"]] = product

    return [products_by_id[product_id] for product_id in product_ids if product_id in products_by_id]


async def _woo_fetch_products_by_ids_store_api(body: WooImportRequest) -> list[dict]:
    product_ids = list(dict.fromkeys(body.product_ids))
    if not product_ids:
        return []

    products_by_id: dict[int, dict] = {}
    async with httpx.AsyncClient(timeout=_WOO_TIMEOUT) as client:
        resolved_endpoint: Optional[str] = None
        for chunk in _chunked(product_ids, _WOO_API_MAX_PER_PAGE):
            resp, resolved_endpoint = await _woo_get_public_store_products(
                client,
                site_url=body.credentials.site_url,
                params={
                    "include": ",".join(str(product_id) for product_id in chunk),
                    "per_page": len(chunk),
                },
                preferred_endpoint=resolved_endpoint,
            )
            for product in resp.json():
                normalized = _woo_store_product_to_woo_shape(product)
                products_by_id[normalized["id"]] = normalized

    return [products_by_id[product_id] for product_id in product_ids if product_id in products_by_id]


async def _woo_fetch_products_by_ids(body: WooImportRequest) -> list[dict]:
    try:
        return await _woo_fetch_products_by_ids_v3(body)
    except HTTPException as exc:
        primary_detail = _woo_preview_error_detail(exc)
        logger.info(
            "Woo v3 import fetch failed for %s, trying public store API fallback: %s",
            body.credentials.site_url,
            primary_detail,
        )
        try:
            return await _woo_fetch_products_by_ids_store_api(body)
        except HTTPException as fallback_exc:
            fallback_detail = _woo_preview_error_detail(fallback_exc)
            raise HTTPException(
                status_code=502,
                detail=f"{primary_detail} | Public Store API fallback failed. {fallback_detail}",
            ) from fallback_exc


@router.post("/woo/preview", response_model=list[WooProductPreview])
async def woo_preview(body: WooPreviewRequest):
    await _woo_validate_preview_request(body)
    return await _woo_fetch_preview_products(body)


@router.post("/woo/import")
async def woo_import(body: WooImportRequest):
    await _woo_validate_import_request(body)
    products = await _woo_fetch_products_by_ids(body)

    db = await get_db()
    imported = 0
    existing_by_url: dict[str, dict] = {}
    existing_by_name: dict[str, dict] = {}
    try:
        existing_by_url, existing_by_name = await _existing_store_product_maps(db, body.store_id)
        for product in products:
            images = product.get("images", [])
            outcome = await _merge_woo_catalog_item(
                db,
                store_id=body.store_id,
                name=product["name"],
                permalink=product["permalink"],
                image_url=images[0]["src"] if images else None,
                existing_by_url=existing_by_url,
                existing_by_name=existing_by_name,
                create_if_missing=True,
            )
            if outcome == "created":
                imported += 1
        await db.commit()
    finally:
        await db.close()

    return {"imported": imported}


@router.post("/api/preview", response_model=list[ApiProductPreview])
async def api_preview(body: ApiPreviewRequest):
    return await fetch_api_preview_products(body)


@router.post("/api/import")
async def api_import(body: ApiImportRequest):
    return await import_api_products(body)


def _to_cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()

def _normalize_headers(raw_headers: list[str]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, header in enumerate(raw_headers):
        base = header.strip() or f"Column {index + 1}"
        count = seen.get(base, 0)
        seen[base] = count + 1
        headers.append(base if count == 0 else f"{base} ({count + 1})")
    return headers


def _tabular_preview(rows: list[list[str]], *, file_format: str, source_name: Optional[str] = None) -> FileImportPreviewResponse:
    nonempty_rows = [row for row in rows if any(cell.strip() for cell in row)]
    if not nonempty_rows:
        raise HTTPException(400, "No usable rows found in the file")

    width = max(len(row) for row in nonempty_rows)
    padded = [row + [""] * (width - len(row)) for row in nonempty_rows]
    columns = _normalize_headers(padded[0])
    data_rows = padded[1:]

    return FileImportPreviewResponse(
        columns=columns,
        sample_rows=data_rows[:10],
        rows=data_rows,
        total_rows=len(data_rows),
        file_format=file_format,
        source_name=source_name,
    )


def _csv_preview(content: bytes) -> FileImportPreviewResponse:
    text: Optional[str] = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise HTTPException(400, "Could not decode CSV file")

    rows = [[_to_cell_str(cell) for cell in row] for row in csv.reader(io.StringIO(text))]
    return _tabular_preview(rows, file_format="csv")


def _excel_preview(content: bytes) -> FileImportPreviewResponse:
    if xlrd is not None:
        try:
            workbook = xlrd.open_workbook(file_contents=content)
        except Exception:
            workbook = None
        else:
            best_sheet_name: Optional[str] = None
            best_rows: list[list[str]] = []
            for worksheet in workbook.sheets():
                rows = [
                    [_to_cell_str(worksheet.cell_value(row_index, col_index)) for col_index in range(worksheet.ncols)]
                    for row_index in range(worksheet.nrows)
                ]
                usable_rows = [row for row in rows if any(cell for cell in row)]
                if len(usable_rows) > len(best_rows):
                    best_rows = usable_rows
                    best_sheet_name = worksheet.name
            if best_rows:
                return _tabular_preview(
                    best_rows,
                    file_format="excel",
                    source_name=best_sheet_name,
                )

    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(content),
            data_only=True,
            read_only=True,
        )
    except Exception as exc:
        if xlrd is None:
            raise HTTPException(
                400,
                f"Could not read Excel file: {exc}. Install xlrd for legacy .xls support.",
            )
        raise HTTPException(400, f"Could not read Excel file: {exc}")

    best_sheet_name: Optional[str] = None
    best_rows: list[list[str]] = []
    try:
        for worksheet in workbook.worksheets:
            rows = [
                [_to_cell_str(cell) for cell in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            usable_rows = [row for row in rows if any(cell for cell in row)]
            if len(usable_rows) > len(best_rows):
                best_rows = usable_rows
                best_sheet_name = worksheet.title
    finally:
        workbook.close()

    if not best_rows:
        raise HTTPException(400, "No usable rows found in the spreadsheet")
    return _tabular_preview(best_rows, file_format="excel", source_name=best_sheet_name)


_PDF_TABLE_STRATEGIES = [
    {
        "vertical_strategy": "lines",
        "horizontal_strategy": "lines",
        "snap_tolerance": 4,
        "join_tolerance": 4,
        "intersection_tolerance": 4,
    },
    {
        "vertical_strategy": "lines_strict",
        "horizontal_strategy": "lines_strict",
        "snap_tolerance": 4,
        "join_tolerance": 4,
        "intersection_tolerance": 4,
    },
    {
        "vertical_strategy": "text",
        "horizontal_strategy": "lines",
        "snap_tolerance": 5,
        "join_tolerance": 5,
        "intersection_tolerance": 5,
        "min_words_vertical": 1,
    },
    {
        "vertical_strategy": "lines",
        "horizontal_strategy": "text",
        "snap_tolerance": 5,
        "join_tolerance": 5,
        "intersection_tolerance": 5,
        "min_words_horizontal": 1,
    },
    {
        "vertical_strategy": "text",
        "horizontal_strategy": "text",
        "text_x_tolerance": 3,
        "text_y_tolerance": 3,
        "min_words_vertical": 1,
        "min_words_horizontal": 1,
    },
]


def _normalize_pdf_table(raw_table) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in raw_table or []:
        normalized = [
            _to_cell_str(cell).replace("\n", " ").strip()
            for cell in (row or [])
        ]
        if any(cell for cell in normalized):
            rows.append(normalized)
    return rows


def _score_pdf_table(rows: list[list[str]]) -> int:
    if not rows:
        return -1
    width = max(len(row) for row in rows)
    if width < 2:
        return -1

    filled_cells = sum(1 for row in rows for cell in row if cell)
    header = rows[0]
    header_alpha_cells = sum(1 for cell in header if re.search(r"[A-Za-z]", cell))
    header_bonus = header_alpha_cells * 3
    data_bonus = 6 if len(rows) > 2 else 0
    consistency_penalty = sum(abs(len(row) - width) for row in rows)
    header_numeric_penalty = sum(
        25 for cell in header if cell and re.fullmatch(r"[\d.,]+", cell)
    )
    header_blank_penalty = sum(12 for cell in header if not cell)
    fragment_penalty = sum(
        6
        for row in rows[: min(len(rows), 8)]
        for cell in row
        if cell and len(cell) <= 3
    )
    width_penalty = max(0, width - 6) * 50

    return (
        len(rows) * 10
        + filled_cells
        + header_bonus
        + data_bonus
        - consistency_penalty
        - header_numeric_penalty
        - header_blank_penalty
        - fragment_penalty
        - width_penalty
    )


def _pdf_table_key(rows: list[list[str]]) -> tuple:
    width = max(len(row) for row in rows)
    padded_header = tuple(
        cell.strip().lower() for cell in (rows[0] + [""] * width)[:width]
    )
    return width, padded_header


def _extract_pdf_text_rows(page) -> list[list[str]]:
    words = page.extract_words(
        keep_blank_chars=False,
        use_text_flow=True,
        x_tolerance=1,
        y_tolerance=3,
    ) or []
    if not words:
        return []

    lines: list[list[dict]] = []
    current_line: list[dict] = []
    current_top: Optional[float] = None
    row_tolerance = 7.5

    for word in sorted(words, key=lambda item: (item["top"], item["x0"])):
        top = float(word["top"])
        if current_line and current_top is not None and abs(top - current_top) > row_tolerance:
            lines.append(current_line)
            current_line = []
            current_top = None
        current_line.append(word)
        if current_top is None:
            current_top = top
        else:
            current_top = (current_top + top) / 2

    if current_line:
        lines.append(current_line)

    rows: list[list[str]] = []
    min_gap = max(18.0, float(page.width) * 0.03)

    for line in lines:
        cells: list[str] = []
        current_cell = ""
        previous_x1: Optional[float] = None

        for word in sorted(line, key=lambda item: item["x0"]):
            x0 = float(word["x0"])
            x1 = float(word["x1"])
            text = _to_cell_str(word["text"])
            if not text:
                continue

            if previous_x1 is not None and x0 - previous_x1 >= min_gap:
                if current_cell.strip():
                    cells.append(current_cell.strip())
                current_cell = text
            else:
                current_cell = f"{current_cell} {text}".strip()

            previous_x1 = x1

        if current_cell.strip():
            cells.append(current_cell.strip())

        if len(cells) >= 2:
            rows.append(cells)

    return rows


def _merge_pdf_tables(candidates: list[dict]) -> list[list[str]]:
    merged: list[list[str]] = []
    seen_headers: set[tuple] = set()

    for candidate in sorted(candidates, key=lambda item: (item["page"], item["order"])):
        rows = candidate["rows"]
        header_key = _pdf_table_key(rows)
        start_index = 1 if header_key in seen_headers and len(rows) > 1 else 0

        for row in rows[start_index:]:
            if merged and row == merged[-1]:
                continue
            merged.append(row)

        seen_headers.add(header_key)

    return merged


def _pdf_preview(content: bytes) -> FileImportPreviewResponse:
    candidates: list[dict] = []
    fallback_rows: list[list[str]] = []

    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page_index, page in enumerate(pdf.pages):
                page_best: Optional[dict] = None
                for settings in _PDF_TABLE_STRATEGIES:
                    try:
                        raw_tables = page.extract_tables(table_settings=settings) or []
                    except Exception:
                        continue

                    for table_index, table in enumerate(raw_tables):
                        rows = _normalize_pdf_table(table)
                        score = _score_pdf_table(rows)
                        if score < 0:
                            continue

                        candidate = {
                            "page": page_index,
                            "order": table_index,
                            "rows": rows,
                            "score": score,
                        }
                        if (
                            page_best is None
                            or score > page_best["score"]
                            or (
                                score == page_best["score"]
                                and len(rows) > len(page_best["rows"])
                            )
                        ):
                            page_best = candidate

                text_rows = _extract_pdf_text_rows(page)
                text_score = _score_pdf_table(text_rows)
                if text_score >= 0:
                    text_candidate = {
                        "page": page_index,
                        "order": 10_000,
                        "rows": text_rows,
                        "score": text_score,
                    }
                    if (
                        page_best is None
                        or text_candidate["score"] > page_best["score"]
                        or (
                            text_candidate["score"] == page_best["score"]
                            and len(text_candidate["rows"]) > len(page_best["rows"])
                        )
                    ):
                        page_best = text_candidate

                if page_best:
                    candidates.append(page_best)
                else:
                    fallback_rows.extend(text_rows)
    except Exception as exc:
        raise HTTPException(400, f"Could not read PDF file: {exc}")

    rows = _merge_pdf_tables(candidates) if candidates else fallback_rows
    return _tabular_preview(rows, file_format="pdf")


def _parse_import_price(value: Optional[float | str]) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value) if float(value) > 0 else None

    cleaned = re.sub(r"[^\d.,]", "", value.strip()).strip(".,")
    if not cleaned:
        return None

    if re.match(r"^\d{1,3}(\.\d{3})+(,\d+)?$", cleaned):
        cleaned = cleaned.replace(".", "").replace(",", ".")
    elif re.match(r"^\d{1,3}(,\d{3})+(\.\d+)?$", cleaned):
        cleaned = cleaned.replace(",", "")
    elif cleaned.count(",") and not cleaned.count("."):
        cleaned = cleaned.replace(",", "." if re.match(r"^\d+,\d{1,2}$", cleaned) else "")
    elif cleaned.count(".") and not cleaned.count(","):
        if re.match(r"^\d{1,3}(\.\d{3})+$", cleaned):
            cleaned = cleaned.replace(".", "")
    else:
        if cleaned.rfind(",") > cleaned.rfind("."):
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")

    try:
        parsed = float(cleaned)
    except ValueError:
        return None
    return parsed if parsed > 0 else None


def _normalize_import_prices(
    price: Optional[float],
    discounted_price: Optional[float],
) -> tuple[Optional[float], Optional[float]]:
    if price is None and discounted_price is not None:
        return discounted_price, None
    if price is None or discounted_price is None:
        return price, discounted_price
    if discounted_price >= price:
        return price, None
    return price, discounted_price


def _normalize_import_stock(status: Optional[str]) -> str:
    text = (status or "").strip().lower()
    if not text:
        return "unknown"
    if any(token in text for token in ("call for price", "call us", "call for availability", "ask for price", "price on request")):
        return "call_for_price"
    if any(token in text for token in ("contact for price", "contact us", "inquire", "inquiry", "whatsapp us", "message us")):
        return "contact_store"
    if any(token in text for token in ("discontinued", "end of life", "eol")):
        return "discontinued"
    if any(token in text for token in ("preorder", "pre-order", "coming soon")):
        return "preorder"
    if any(token in text for token in ("backorder", "back-order")):
        return "backorder"
    if any(token in text for token in ("out of stock", "sold out", "unavailable", "no stock")):
        return "out_of_stock"
    if any(token in text for token in ("limited", "low stock", "few left")):
        return "limited"
    if any(token in text for token in ("in stock", "available", "instock")):
        return "in_stock"
    return "unknown"


def _effective_price(price: Optional[float], discounted_price: Optional[float]) -> Optional[float]:
    return discounted_price if discounted_price is not None else price


def _compute_file_change(previous: Optional[dict], row: FileImportRow) -> tuple[str, Optional[float], Optional[float], Optional[str]]:
    current_effective = _effective_price(row.price, row.discounted_price)
    current_stock = row.stock_status or "unknown"
    if previous is None:
        return "new", None, None, None

    previous_effective = _effective_price(previous.get("price"), previous.get("discounted_price"))
    previous_stock = previous.get("stock_status") or "unknown"

    price_changed = False
    price_delta: Optional[float] = None
    if previous_effective is not None and current_effective is not None:
        delta = current_effective - previous_effective
        if abs(delta) > 0.001:
            price_delta = round(delta, 2)
            price_changed = True
    elif previous_effective is None and current_effective is not None:
        price_changed = True

    stock_changed = previous_stock != current_stock

    if price_changed and stock_changed:
        return "price_and_stock", price_delta, previous_effective, previous_stock
    if price_changed:
        if price_delta is not None:
            return ("price_up" if price_delta > 0 else "price_down"), price_delta, previous_effective, previous_stock
        return "price_up", price_delta, previous_effective, previous_stock
    if stock_changed:
        if current_stock == "out_of_stock":
            return "stock_out", price_delta, previous_effective, previous_stock
        if current_stock in ("in_stock", "limited"):
            return "stock_in", price_delta, previous_effective, previous_stock
        return "stock_change", price_delta, previous_effective, previous_stock
    return "no_change", price_delta, previous_effective, previous_stock


def _placeholder_product_url(store_id: int, row: FileImportRow) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", (row.name or "").strip().lower()).strip("-")
    if not base:
        base = "item"
    suffix = row.source_index if row.source_index is not None else base
    return f"import://store-{store_id}/{quote(base)}-{suffix}"


def _stock_hint_from_row(row: FileImportRow) -> str:
    for candidate in (
        row.stock_status,
        row.price if isinstance(row.price, str) else None,
        row.discounted_price if isinstance(row.discounted_price, str) else None,
        row.dealer_price if isinstance(row.dealer_price, str) else None,
    ):
        status = _normalize_import_stock(candidate)
        if status != "unknown":
            return status
    return "unknown"


@router.post("/file/preview", response_model=FileImportPreviewResponse)
async def file_preview(file: UploadFile = File(...)):
    content = await file.read()
    filename = (file.filename or "").lower()

    if filename.endswith(".csv"):
        return _csv_preview(content)

    if filename.endswith((".xlsx", ".xlsm", ".xls")):
        return _excel_preview(content)

    if filename.endswith(".pdf"):
        return _pdf_preview(content)

    raise HTTPException(400, "Unsupported format - use .csv, .xlsx, or .pdf")


async def _find_import_target(db, *, canonical_name: str, store_id: int):
    rows = await db.execute_fetchall(
        """SELECT
               p.id AS product_id,
               p.image_url,
               p.category_id,
               ps.id AS product_store_id,
               ps.product_url,
               ps.is_active
           FROM products p
           JOIN product_stores ps ON ps.product_id = p.id
           WHERE ps.store_id = ?
             AND lower(trim(p.canonical_name)) = lower(trim(?))
           ORDER BY p.id
           LIMIT 1""",
        (store_id, canonical_name),
    )
    return rows[0] if rows else None


async def _find_import_target_by_product(db, *, product_id: int, store_id: int):
    rows = await db.execute_fetchall(
        """SELECT
               p.id AS product_id,
               p.canonical_name,
               p.image_url,
               p.category_id,
               ps.id AS product_store_id,
               ps.store_id,
               ps.product_url,
               ps.is_active
           FROM products p
           LEFT JOIN product_stores ps
             ON ps.product_id = p.id AND ps.store_id = ?
           WHERE p.id = ?
           LIMIT 1""",
        (store_id, product_id),
    )
    return rows[0] if rows else None


async def _find_import_target_by_store_link(db, *, product_store_id: int):
    rows = await db.execute_fetchall(
        """SELECT
               p.id AS product_id,
               p.canonical_name,
               p.image_url,
               p.category_id,
               ps.id AS product_store_id,
               ps.store_id,
               ps.product_url,
               ps.is_active
           FROM product_stores ps
           JOIN products p ON p.id = ps.product_id
           WHERE ps.id = ?
           LIMIT 1""",
        (product_store_id,),
    )
    return rows[0] if rows else None


async def _latest_price_log(db, product_store_id: int):
    rows = await db.execute_fetchall(
        """SELECT price, discounted_price, stock_status
           FROM price_logs
           WHERE product_store_id = ?
           ORDER BY scraped_at DESC
           LIMIT 1""",
        (product_store_id,),
    )
    return dict(rows[0]) if rows else None


async def _upsert_import_target(
    db,
    *,
    row: FileImportRow,
    store_id: int,
    category_id: Optional[int],
    create_missing_products: bool,
) -> tuple[Optional[int], Optional[int], str]:
    provided_url = (row.url or "").strip()
    has_real_url = bool(provided_url)
    entered_name = row.name.strip()
    explicit_target = None
    if row.product_store_id is not None:
        explicit_target = await _find_import_target_by_store_link(
            db,
            product_store_id=row.product_store_id,
        )
        if explicit_target is None:
            raise HTTPException(400, f"Product store {row.product_store_id} was not found")
        if explicit_target["store_id"] != store_id:
            raise HTTPException(400, "Selected product store does not belong to the chosen store")
    elif row.product_id is not None:
        explicit_target = await _find_import_target_by_product(
            db,
            product_id=row.product_id,
            store_id=store_id,
        )
        if explicit_target is None:
            raise HTTPException(400, f"Product {row.product_id} was not found")

    canonical_name = (
        (explicit_target["canonical_name"] if explicit_target is not None else None)
        or entered_name
    )
    if not canonical_name:
        raise HTTPException(400, "Imported rows must include a product name or an existing product target")

    row.name = entered_name or canonical_name
    resolved_url = provided_url or _placeholder_product_url(store_id, row)
    target = explicit_target
    if target is None:
        target = await _find_import_target(db, canonical_name=canonical_name, store_id=store_id)

    if target is None:
        if not create_missing_products:
            return None, None, "skipped"
        product_id, product_store_id = await create_product_with_store(
            db,
            canonical_name=canonical_name,
            image_url=(row.image_url or "").strip() or None,
            category_id=category_id,
            store_id=store_id,
            product_url=resolved_url,
            is_active=has_real_url,
        )
        return product_id, product_store_id, "created"

    product_id = target["product_id"]
    product_store_id = target["product_store_id"]

    if product_store_id is None:
        product_store_id = await create_store_link(
            db,
            product_id=product_id,
            store_id=store_id,
            product_url=resolved_url,
            is_active=has_real_url,
        )
        return product_id, product_store_id, "updated"

    if row.image_url and not target["image_url"]:
        await db.execute(
            "UPDATE products SET image_url = ? WHERE id = ?",
            (row.image_url, product_id),
        )
    if category_id is not None and target["category_id"] is None:
        await db.execute(
            "UPDATE products SET category_id = ? WHERE id = ?",
            (category_id, product_id),
        )

    should_replace_url = (
        has_real_url
        and (
            not target["product_url"]
            or str(target["product_url"]).startswith("import://")
            or str(target["product_url"]).strip() != provided_url
        )
    )
    if should_replace_url:
        await db.execute(
            """UPDATE product_stores
               SET product_url = ?, is_active = 1
               WHERE id = ?""",
            (provided_url, product_store_id),
        )

    return product_id, product_store_id, "updated"


async def _insert_import_snapshot(db, *, product_store_id: int, row: FileImportRow) -> None:
    has_snapshot = (
        row.price is not None
        or row.discounted_price is not None
        or row.dealer_price is not None
        or row.warranty not in (None, "")
        or row.stock_status not in (None, "", "unknown")
    )
    if not has_snapshot:
        return

    previous = await _latest_price_log(db, product_store_id)
    change_type, price_delta, previous_price, previous_stock_status = _compute_file_change(previous, row)

    await db.execute(
        """INSERT INTO price_logs
           (product_store_id, title, price, discounted_price, dealer_price, warranty,
            stock_status, raw_price, raw_discounted, raw_dealer,
            change_type, price_delta, previous_price, previous_stock_status)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (
            product_store_id,
            row.name,
            row.price,
            row.discounted_price,
            row.dealer_price,
            row.warranty,
            row.stock_status or "unknown",
            str(row.price) if row.price is not None else None,
            str(row.discounted_price) if row.discounted_price is not None else None,
            str(row.dealer_price) if row.dealer_price is not None else None,
            change_type,
            price_delta,
            previous_price,
            previous_stock_status,
        ),
    )


async def _set_import_fallback(
    db,
    *,
    product_store_id: int,
    dealer_price: Any = _KEEP_IMPORT_VALUE,
    warranty: Any = _KEEP_IMPORT_VALUE,
) -> None:
    if dealer_price is _KEEP_IMPORT_VALUE and warranty is _KEEP_IMPORT_VALUE:
        return

    rows = await db.execute_fetchall(
        """SELECT
               manual_price,
               manual_stock_status,
               manual_warranty,
               imported_dealer_price,
               imported_warranty,
               override_note
           FROM product_store_overrides
           WHERE product_store_id = ?
           LIMIT 1""",
        (product_store_id,),
    )
    current = rows[0] if rows else None

    normalized_warranty = (
        (warranty or "").strip() or None
        if warranty is not _KEEP_IMPORT_VALUE
        else (
            (current["imported_warranty"] or "").strip() or None
            if current
            else None
        )
    )
    next_dealer_price = (
        current["imported_dealer_price"]
        if dealer_price is _KEEP_IMPORT_VALUE and current
        else (None if dealer_price is _KEEP_IMPORT_VALUE else dealer_price)
    )

    if current is None and next_dealer_price is None and normalized_warranty is None:
        return

    await db.execute(
        """INSERT INTO product_store_overrides
               (product_store_id, imported_dealer_price, imported_warranty)
           VALUES (?, ?, ?)
           ON CONFLICT(product_store_id) DO UPDATE SET
               imported_dealer_price = excluded.imported_dealer_price,
               imported_warranty = excluded.imported_warranty,
               updated_at = CURRENT_TIMESTAMP""",
        (
            product_store_id,
            next_dealer_price,
            normalized_warranty,
        ),
    )

    await db.execute(
        """DELETE FROM product_store_overrides
           WHERE product_store_id = ?
             AND manual_price IS NULL
             AND manual_stock_status IS NULL
             AND manual_warranty IS NULL
             AND imported_dealer_price IS NULL
             AND imported_warranty IS NULL
             AND override_note IS NULL""",
        (product_store_id,),
    )


def _normalize_seen_product_store_ids(values: list[int]) -> set[int]:
    return {
        int(value)
        for value in values
        if isinstance(value, int) and value > 0
    }


async def _mark_missing_full_catalog_dealer_unavailable(
    db,
    *,
    store_id: int,
    seen_product_store_ids: set[int],
) -> int:
    params: list[Any] = [store_id]
    query = """
        SELECT pso.product_store_id
        FROM product_store_overrides pso
        JOIN product_stores ps ON ps.id = pso.product_store_id
        WHERE ps.store_id = ?
          AND (
              pso.imported_dealer_price IS NOT NULL
              OR NULLIF(TRIM(pso.imported_warranty), '') IS NOT NULL
          )
    """
    if seen_product_store_ids:
        placeholders = ", ".join("?" for _ in seen_product_store_ids)
        query += f" AND pso.product_store_id NOT IN ({placeholders})"
        params.extend(sorted(seen_product_store_ids))

    rows = await db.execute_fetchall(query, tuple(params))
    missing_ids = [int(row["product_store_id"]) for row in rows]
    if not missing_ids:
        return 0

    placeholders = ", ".join("?" for _ in missing_ids)
    await db.execute(
        f"""UPDATE product_store_overrides
            SET imported_dealer_price = NULL,
                imported_warranty = NULL,
                updated_at = CURRENT_TIMESTAMP
            WHERE product_store_id IN ({placeholders})""",
        tuple(missing_ids),
    )
    await db.execute(
        f"""DELETE FROM product_store_overrides
            WHERE product_store_id IN ({placeholders})
              AND manual_price IS NULL
              AND manual_stock_status IS NULL
              AND manual_warranty IS NULL
              AND imported_dealer_price IS NULL
              AND imported_warranty IS NULL
              AND override_note IS NULL""",
        tuple(missing_ids),
    )
    return len(missing_ids)


@router.post("/file/commit", response_model=FileImportCommitResponse)
async def file_commit(body: FileImportCommitRequest):
    db = await get_db()
    imported = 0
    created = 0
    updated = 0
    skipped = 0
    dealer_unavailable_marked = 0
    processed_product_store_ids: set[int] = set()
    try:
        for row in body.rows:
            if not row.name.strip() and row.product_id is None and row.product_store_id is None:
                continue

            row.warranty = (row.warranty or "").strip() or None
            hinted_stock = _stock_hint_from_row(row)
            row.price = _parse_import_price(row.price)
            row.discounted_price = _parse_import_price(row.discounted_price)
            row.dealer_price = _parse_import_price(row.dealer_price)
            row.price, row.discounted_price = _normalize_import_prices(
                row.price,
                row.discounted_price,
            )
            row.stock_status = _normalize_import_stock(row.stock_status)
            if row.stock_status == "unknown":
                row.stock_status = hinted_stock

            _, product_store_id, action = await _upsert_import_target(
                db,
                row=row,
                store_id=body.store_id,
                category_id=body.category_id,
                create_missing_products=body.create_missing_products,
            )
            if action == "skipped" or product_store_id is None:
                skipped += 1
                continue
            processed_product_store_ids.add(product_store_id)

            await _set_import_fallback(
                db,
                product_store_id=product_store_id,
                dealer_price=(
                    row.dealer_price
                    if row.dealer_price is not None
                    else (
                        None
                        if body.dealer_column_mapped
                        and body.missing_rows_policy == "mark_dealer_unavailable"
                        else _KEEP_IMPORT_VALUE
                    )
                ),
                warranty=(
                    row.warranty
                    if row.warranty is not None
                    else (
                        None
                        if body.warranty_column_mapped
                        and body.missing_rows_policy == "mark_dealer_unavailable"
                        else _KEEP_IMPORT_VALUE
                    )
                ),
            )
            await _insert_import_snapshot(db, product_store_id=product_store_id, row=row)
            imported += 1
            if action == "created":
                created += 1
            else:
                updated += 1

        if body.missing_rows_policy == "mark_dealer_unavailable":
            seen_product_store_ids = _normalize_seen_product_store_ids(
                body.full_catalog_seen_product_store_ids
            )
            seen_product_store_ids.update(processed_product_store_ids)
            dealer_unavailable_marked = (
                await _mark_missing_full_catalog_dealer_unavailable(
                    db,
                    store_id=body.store_id,
                    seen_product_store_ids=seen_product_store_ids,
                )
            )
        await db.commit()
    finally:
        await db.close()

    return FileImportCommitResponse(
        imported=imported,
        created=created,
        updated=updated,
        skipped=skipped,
        dealer_unavailable_marked=dealer_unavailable_marked,
    )
